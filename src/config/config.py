import json
import warnings
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from src.utils.path_manager import get_resource


class Config:
    def __init__(self):
        self.signal_map     = None       # vbRcSignals sheet
        self.event_kpi_list = None       # event KPI sheet
        self.cycle_kpi_list = None       # cycle KPI sheet
        self.graph_spec     = None       # graphSpec sheet
        self.line_colors    = None       # lineColors sheet
        self.marker_shapes  = None       # markerShapes sheet
        self.calibratables  = {}         # calibratables (optional)
        self.params         = None       # params sheet 
        self.param_types    = {}         # keep parameter type metadata

    @classmethod
    def from_json(cls, json_config_path):
        """Create Config object from JSON file (supports both as_long & as_lat)."""
        cfg = cls()
        config_struct = cls._load_config(json_config_path)

        # =====================================================
        # 1️⃣ Locate KPI workbook (for KPI sheets and vbRcSignals)
        # =====================================================
        kpi_section_key = next(
            (k for k in config_struct.keys() if k.lower().startswith("kpias")),
            None,
        )
        if not kpi_section_key:
            raise ValueError("No KPI section found in config (expected 'KpiAsLong' or 'KpiAsLat').")

        spec_cfg = config_struct[kpi_section_key]
        spec_path = get_resource(f"config/{spec_cfg['FilePath']}")

        # =====================================================
        # 2️⃣ Load vbRcSignals from KPI workbook
        # =====================================================
        sig_data = cls._load_signal_map_kpi_plot_spec(spec_path, ["vbRcSignals"])
        sig_data = {k.lower(): v for k, v in sig_data.items()}
        cfg.signal_map = sig_data.get("vbrcsignals")

        if cfg.signal_map is None:
            raise ValueError(
                f"Signal sheet 'vbRcSignals' not found in {Path(spec_path).name}. "
                "Ensure the KPI workbook includes a 'vbRcSignals' sheet."
            )
        print(f"✅ Signal map loaded from {Path(spec_path).name} → sheet 'vbRcSignals'")
        cfg.signal_map = cls._normalize_columns(cfg.signal_map)

        # =====================================================
        # 3️⃣ Load KPI/PlotSpec-related sheets (auto-detect Long/Lat/etc.)
        # =====================================================
        sheet_list = spec_cfg["Sheets"]

        ## --- Define the name for the KPI excel export ---
        cfg.kpi_excel_path = Path(spec_path)
        cfg.kpi_excel_name = cfg.kpi_excel_path.stem.lower()        # e.g. "kpi_as_lat"
        cfg.domain_name = "as_long" if "long" in cfg.kpi_excel_name else (
            "as_lat" if "lat" in cfg.kpi_excel_name else "as_unknown"
        )
        cfg.kpi_result_filename = f"{cfg.domain_name}_kpi_results.xlsx"

        print(f"📘 Loading KPI section: '{kpi_section_key}' → {Path(spec_path).name}")

        spec_data  = cls._load_signal_map_kpi_plot_spec(spec_path, sheet_list)
        sheet_map  = {k.lower(): cls._normalize_columns(v) for k, v in spec_data.items()}

        cfg.graph_spec     = sheet_map.get("graphspec")
        cfg.line_colors    = sheet_map.get("linecolors")
        cfg.marker_shapes  = sheet_map.get("markershapes")
        cfg.event_kpi_list = sheet_map.get("kpi")
        cfg.params         = sheet_map.get("params")

        cfg.cycle_kpi_list = sheet_map.get("cyclekpi")
        if getattr(cfg.cycle_kpi_list, "empty", True):
            warnings.warn("⚠️ No cycleKPI/overallKPI sheet found in config workbook.")


        # =====================================================
        # 4️⃣ Parse params sheet into dict with type awareness
        # =====================================================
        if cfg.params is not None and not cfg.params.empty:
            try:
                cfg.params, cfg.param_types = cls._parse_params_sheet(cfg.params)
                print(f"⚙️ Loaded {len(cfg.params)} parameters from 'params' sheet.")
                print("   ➝ Keys:", ", ".join(list(cfg.params.keys())[:6]), "...")
            except Exception as e:
                warnings.warn(f"⚠️ Failed to parse params sheet: {e}")

        # =====================================================
        # 5️⃣ Normalize and clean line_colors sheet
        # =====================================================
        if cfg.line_colors is not None and not cfg.line_colors.empty:
            rgb_cols = [c for c in cfg.line_colors.columns if c in ["r", "g", "b"]]
            if len(rgb_cols) == 3:
                try:
                    cfg.line_colors = (
                        cfg.line_colors[rgb_cols]
                        .astype(float)
                        .clip(0, 1)
                        .to_numpy()
                        .tolist()
                    )
                    print(f"🎨 Loaded {len(cfg.line_colors)} RGB color entries from lineColors.")
                except Exception as e:
                    warnings.warn(f"⚠️ Failed to parse RGB colors from lineColors: {e}")
            else:
                warnings.warn("⚠️ No valid R,G,B columns found in lineColors sheet.")

        # =====================================================
        # 6️⃣ Load calibratables from calParam in KPI workbook (Optional)
        # =====================================================
        cal_defs = spec_cfg.get("Calibratables", None)

        if cal_defs:
            print(f"📗 Loading calibratables from '{cfg.kpi_excel_path.name}' (sheet 'calParam')")

            # wrap in dict so _load_calibratables understands it
            sheet_map = {"calParam": cal_defs}

            cfg.calibratables = cls._load_calibratables(
                cfg.kpi_excel_path,
                sheet_map
            )

            cfg._apply_calibration_scaling()
        else:
            cfg.calibratables = {}
            print("⚙️ No Calibratables defined under KPI section.")


        return cfg


    # =====================================================
    # Helper functions
    # =====================================================
    @staticmethod
    def _load_config(file_path):
        """
        Load and validate the JSON configuration file.

        Automatically detects KPI section (e.g. KpiAsLong, KpiAsLat, etc.).
        Calibration section is optional for lateral (as_lat) configs.
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Config file not found: {file_path}")

        with open(file_path, "r", encoding="utf-8") as f:
            params = json.load(f)

        if "SignalMap" in params:
            warnings.warn("⚠️ SignalMap section is no longer used; remove it from the config JSON.")


        # --- Detect KPI section automatically (any key starting with 'KpiAs') ---
        kpi_section_key = next(
            (k for k in params.keys() if k.lower().startswith("kpias")),
            None
        )

        if not kpi_section_key:
            raise ValueError(
                "Missing KPI section (expected key starting with 'KpiAs', "
                "e.g. 'KpiAsLong' or 'KpiAsLat')."
            )

        kpi_section = params[kpi_section_key]
        if "FilePath" not in kpi_section:
            raise ValueError(f"Missing {kpi_section_key}.FilePath in config.")
        if "Sheets" not in kpi_section:
            raise ValueError(f"{kpi_section_key}.Sheets must be defined in config.")

        print(f"📘 Detected KPI section: '{kpi_section_key}'")

        return params

    @staticmethod
    def _normalize_columns(df):
        if df is None:
            return None
        df.columns = df.columns.str.strip().str.lower()
        return df

    @staticmethod
    def _parse_params_sheet(params_df):
        df = params_df.dropna(subset=["parameter", "value"]).copy()
        df["parameter"] = df["parameter"].astype(str).str.strip().str.lower()

        param_dict = {}
        type_dict = {}

        for _, row in df.iterrows():
            name = row["parameter"]
            value = row["value"]
            ptype = str(row.get("type", "")).strip().lower()

            try:
                if ptype in ("int", "integer"):
                    cast_val = int(float(value))
                elif ptype in ("float", "double", "numeric"):
                    cast_val = float(value)
                elif ptype in ("bool", "boolean"):
                    cast_val = bool(value)
                elif ptype in ("str", "string"):
                    cast_val = str(value)
                else:
                    cast_val = float(value)
            except Exception:
                cast_val = value

            param_dict[name] = cast_val
            type_dict[name] = ptype or type(cast_val).__name__

        return param_dict, type_dict


    @staticmethod
    def _load_signal_map_kpi_plot_spec(file_path, sheet_list):
        """Load multiple sheets with automatic header detection."""
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        excel = pd.ExcelFile(file_path)
        result = {}
        for sheet_name in sheet_list:

            # -----------------------------------------------
            # Skip sheets that do not exist in workbook
            # -----------------------------------------------
            if sheet_name not in excel.sheet_names:
                warnings.warn(f"⚠️ Sheet '{sheet_name}' not found in {file_path.name}, skipping.")
                continue

            try:
                preview = excel.parse(sheet_name=sheet_name, nrows=5, header=None)

                header_row = next(
                    (i for i in range(len(preview)) if preview.iloc[i].notna().sum() >= 3),
                    0,
                )

                df = excel.parse(sheet_name=sheet_name, header=header_row)
                df.columns = df.columns.str.strip().str.lower()

                result[sheet_name] = df
                print(f"✅ Loaded '{sheet_name}' (header at row {header_row+1}) — shape {df.shape}")

            except Exception as e:
                warnings.warn(f'⚠️ Failed to read sheet "{sheet_name}": {e}')

        return result

    @staticmethod
    def _load_calibratables(calib_file, sheet_map):
        file_path = Path(calib_file)
        if not file_path.exists():
            raise FileNotFoundError(f"Calibration file not found: {file_path}")

        wb = load_workbook(file_path, data_only=True)
        calibratables = {}

        for sheet, cal_defs in sheet_map.items():
            if sheet not in wb.sheetnames:
                warnings.warn(f'Sheet "{sheet}" not found in calibration file.')
                continue

            ws = wb[sheet]
            for cal_name, rng in cal_defs.items():
                try:
                    cells = ws[rng]
                    data = [[cell.value for cell in row] for row in cells]
                    df = pd.DataFrame(data).reset_index(drop=True)

                    if df.isna().all().all():
                        calibratables[cal_name] = None
                    elif df.shape[0] == 2:
                        x = [v if v is not None else float("nan") for v in df.iloc[0].tolist()]
                        y = [v if v is not None else float("nan") for v in df.iloc[1].tolist()]
                        calibratables[cal_name] = {"x": x, "y": y}
                    else:
                        calibratables[cal_name] = df
                except Exception as e:
                    warnings.warn(
                        f'Failed to load "{cal_name}" from sheet "{sheet}" range "{rng}": {e}'
                    )
                    calibratables[cal_name] = None

        wb.close()
        return calibratables

    def _apply_calibration_scaling(self):
        """
        Apply scaling or normalization logic to certain calibratables.
        """
        key = "PedalPosProIncrease_Th"
        if key in self.calibratables:
            val = self.calibratables[key]
            if isinstance(val, dict) and "y" in val:
                y_vals = val["y"]
                if all(isinstance(v, (int, float)) for v in y_vals if v is not None):
                    if all(0 <= v <= 1 for v in y_vals):
                        self.calibratables[key]["y"] = [
                            v * 100 if v is not None else None for v in y_vals
                        ]
                        print(f"📏 Scaled '{key}' *100 (0-1 → 0-100).")
